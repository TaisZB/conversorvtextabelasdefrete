from pathlib import Path

from openpyxl import load_workbook
import pandas as pd


class ExcelReader:
    """
    Responsável por abrir e inspecionar planilhas Excel.
    """

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"Arquivo não encontrado: {self.file_path}"
            )

        suffix = self.file_path.suffix.lower()

        if suffix == ".xls":
            self.workbook = pd.ExcelFile(self.file_path)
            self._is_legacy = True
        else:
            self.workbook = load_workbook(
                filename=self.file_path,
                data_only=True
            )
            self._is_legacy = False

    def get_sheet_names(self) -> list[str]:
        """Retorna todas as abas do arquivo."""
        if self._is_legacy:
            return self.workbook.sheet_names
        return self.workbook.sheetnames

    def get_sheet(self, sheet_name=None):
        """Retorna uma aba."""
        if self._is_legacy:
            if sheet_name is None:
                sheet_name = self.workbook.sheet_names[0]
            return pd.read_excel(self.file_path, sheet_name=sheet_name)

        if sheet_name is None:
            return self.workbook[self.workbook.sheetnames[0]]

        return self.workbook[sheet_name]

    def get_dimensions(self, sheet_name=None):
        """Quantidade de linhas e colunas."""

        if self._is_legacy:
            sheet = self.get_sheet(sheet_name)
            return {
                "rows": len(sheet.index),
                "columns": len(sheet.columns)
            }

        sheet = self.get_sheet(sheet_name)

        return {
            "rows": sheet.max_row,
            "columns": sheet.max_column
        }

    def preview(self, sheet_name=None, rows=10):
        """Retorna as primeiras linhas."""

        if self._is_legacy:
            sheet = self.get_sheet(sheet_name)
            return sheet.head(rows).values.tolist()

        sheet = self.get_sheet(sheet_name)

        result = []

        for row in sheet.iter_rows(
            min_row=1,
            max_row=rows,
            values_only=True
        ):
            result.append(list(row))

        return result
